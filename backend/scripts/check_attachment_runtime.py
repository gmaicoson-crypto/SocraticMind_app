"""Deployment smoke check: private sandbox, CJK fonts and real DOC/PPT conversion.

Run from the repository root with .venv/bin/python backend/scripts/check_attachment_runtime.py.
No database connection or model request is made.
"""

import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from pptx import Presentation

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))
os.chdir(BACKEND_DIR)

# Application settings must be imported after selecting the backend .env directory.
from services.attachments.config import isolation_available, libreoffice_path  # noqa: E402
from services.attachments.worker import execute_parser  # noqa: E402


def cjk_fonts_present() -> bool:
    fontconfig = shutil.which("fc-list")
    if fontconfig:
        result = subprocess.run([fontconfig, ":lang=zh", "family"], capture_output=True, timeout=10)
        if result.returncode == 0 and result.stdout.strip():
            return True
    for root in (Path("/System/Library/Fonts"), Path("/Library/Fonts"), Path("/usr/share/fonts")):
        if root.is_dir() and any(
            any(
                name in font.name.lower()
                for name in ("pingfang", "heiti", "songti", "notosanscjk", "notoserifcjk", "wqy")
            )
            for font in root.rglob("*")
        ):
            return True
    return False


def main() -> int:
    executable = libreoffice_path()
    checks = {
        "sandbox": isolation_available(),
        "libreoffice": bool(executable),
        "cjk_fonts": cjk_fonts_present(),
    }
    if not all(checks.values()):
        print(json.dumps({"ok": False, "checks": checks}, ensure_ascii=False, indent=2))
        return 1
    with tempfile.TemporaryDirectory(prefix="attachment-runtime-check-") as temporary:
        directory = Path(temporary).resolve()
        for legacy, modern in (("doc", "docx"), ("ppt", "pptx")):
            source = directory / f"fixture.{modern}"
            if modern == "docx":
                document = Document()
                document.add_paragraph("中文资料与 English content")
                document.save(source)
            else:
                presentation = Presentation()
                slide = presentation.slides.add_slide(presentation.slide_layouts[1])
                slide.shapes.title.text = "中文资料与 English content"
                slide.notes_slide.notes_text_frame.text = "教师备注"
                presentation.save(source)
            profile = directory / f"profile-{legacy}"
            result = subprocess.run(
                [
                    executable,
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--headless",
                    "--convert-to",
                    legacy,
                    "--outdir",
                    str(directory),
                    str(source),
                ],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=45,
                env={
                    "PATH": os.environ.get("PATH", "/usr/bin:/bin"),
                    "HOME": str(directory),
                    "TMPDIR": str(directory),
                },
            )
            old_source = directory / f"fixture.{legacy}"
            output = directory / f"output-{legacy}"
            output.mkdir()
            parsed = (
                execute_parser(old_source, output)
                if result.returncode == 0 and old_source.exists()
                else {}
            )
            checks[legacy] = parsed.get("status") == "ready" and any(
                "中文资料" in block["text"] for block in parsed.get("chunks", [])
            )
    print(json.dumps({"ok": all(checks.values()), "checks": checks}, ensure_ascii=False, indent=2))
    return 0 if all(checks.values()) else 1


if __name__ == "__main__":
    raise SystemExit(main())

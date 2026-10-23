"""Bounded parsers. The worker executes this file in an OS sandbox, never in chat."""

from __future__ import annotations

import json
from pathlib import Path, PurePosixPath
import stat
import subprocess
import sys
import zipfile

if __package__:
    from .config import LIMITS, MAX_BLOCKS, MAX_TEXT_CHARS, ZIP_MAX_BYTES, ZIP_MAX_ENTRIES
else:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from config import LIMITS, MAX_BLOCKS, MAX_TEXT_CHARS, ZIP_MAX_BYTES, ZIP_MAX_ENTRIES


class ParseFailure(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


def fail(message: str, code: str = "ATTACHMENT_PARSE_FAILED") -> None:
    raise ParseFailure(code, message)


def inspect_zip(path: Path, extension: str) -> None:
    """Read only the archive directory before any decompression occurs."""
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > ZIP_MAX_ENTRIES:
            fail("Office 文件内部条目过多")
        total = 0
        names = set()
        for entry in entries:
            name = PurePosixPath(entry.filename)
            if (
                name.is_absolute()
                or ".." in name.parts
                or "\\" in entry.filename
                or stat.S_ISLNK(entry.external_attr >> 16)
            ):
                fail("Office 文件包含不安全的内部路径")
            if entry.flag_bits & 1:
                fail("暂不支持加密文件")
            total += entry.file_size
            if total > ZIP_MAX_BYTES or entry.file_size > max(1, entry.compress_size) * 200:
                fail("Office 文件解压大小或压缩比例超过限制")
            names.add(entry.filename)
        required = "word/document.xml" if extension == ".docx" else "ppt/presentation.xml"
        if required not in names or "[Content_Types].xml" not in names:
            fail("文件内容与扩展名不一致", "ATTACHMENT_TYPE_UNSUPPORTED")
        if any("vbaproject" in name.lower() for name in names):
            fail("不支持包含宏的 Office 文件", "ATTACHMENT_TYPE_UNSUPPORTED")


def validate_file_type(path: Path, extension: str) -> None:
    with path.open("rb") as stream:
        magic = stream.read(16)
    if extension in {".docx", ".pptx"}:
        if not magic.startswith(b"PK"):
            fail("Office 文件签名不匹配", "ATTACHMENT_TYPE_UNSUPPORTED")
        inspect_zip(path, extension)
    elif extension in {".doc", ".ppt"}:
        if not magic.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            fail("旧版 Office 文件签名不匹配", "ATTACHMENT_TYPE_UNSUPPORTED")
        import olefile

        with olefile.OleFileIO(path) as archive:
            names = {"/".join(item) for item in archive.listdir()}
            if any("EncryptedPackage" in name or "EncryptionInfo" in name for name in names):
                fail("暂不支持加密文件")
            expected = "WordDocument" if extension == ".doc" else "PowerPoint Document"
            if expected not in names:
                fail("文件内容与扩展名不一致", "ATTACHMENT_TYPE_UNSUPPORTED")
    elif extension == ".pdf":
        if not magic.startswith(b"%PDF-"):
            fail("PDF 文件签名不匹配", "ATTACHMENT_TYPE_UNSUPPORTED")
    else:
        from PIL import Image

        with Image.open(path) as image:
            expected = {".jpg": "JPEG", ".jpeg": "JPEG", ".png": "PNG", ".webp": "WEBP"}.get(
                extension
            )
            if expected != image.format or getattr(image, "is_animated", False):
                fail("不支持该图片格式或动画图片", "ATTACHMENT_TYPE_UNSUPPORTED")
            if image.width * image.height > LIMITS["max_image_pixels"]:
                fail("图片像素超过 4000 万限制", "ATTACHMENT_TOO_LARGE")
            image.verify()


class Blocks:
    def __init__(self):
        self.items: list[dict] = []
        self.characters = 0

    def add(
        self,
        text: str,
        locator: dict,
        block_type: str = "paragraph",
        title_path: list | None = None,
    ) -> None:
        text = text.strip()
        if not text:
            return
        self.characters += len(text)
        if self.characters > MAX_TEXT_CHARS or len(self.items) >= MAX_BLOCKS:
            fail("文档文本或结构数量超过限制，请拆分文件", "ATTACHMENT_TOO_LARGE")
        # Preserve ordinary paragraphs/table rows; split only unusually long blocks.
        for offset in range(0, len(text), 3200):
            self.items.append(
                {
                    "chunk_id": f"c{len(self.items) + 1}",
                    "position": len(self.items),
                    "text": text[offset : offset + 3200],
                    "locator": locator,
                    "block_type": block_type,
                    "title_path": list(title_path or []),
                    "asset_refs": [],
                }
            )


def parse_docx(path: Path, blocks: Blocks) -> dict:
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    document = Document(path)
    headings: list[str] = []
    counter = {"paragraph": 0, "table": 0}

    def visit_table(table, number, depth):
        if depth > 12:
            fail("Word 嵌套表格层级超过限制")
        seen = set()
        for row in table.rows:
            cells = []
            nested_tables = []
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                cells.append("\n".join(p.text for p in cell.paragraphs))
                nested_tables.extend(cell.tables)
            blocks.add(
                " | ".join(cells),
                {"table": number, "section": " / ".join(headings)[:1000]},
                "table",
                headings,
            )
            for nested in nested_tables:
                visit_table(nested, number, depth + 1)

    def visit(container):
        for item in container.iter_inner_content():
            if isinstance(item, Paragraph):
                counter["paragraph"] += 1
                style = item.style.name if item.style else ""
                if style.startswith("Heading ") and style[8:].isdigit():
                    level = min(int(style[8:]), 8)
                    headings[:] = headings[: level - 1] + [item.text[:200]]
                blocks.add(
                    item.text,
                    {"paragraph": counter["paragraph"], "section": " / ".join(headings)[:1000]},
                    "heading" if style.startswith("Heading") else "paragraph",
                    headings,
                )
            elif isinstance(item, Table):
                counter["table"] += 1
                visit_table(item, counter["table"], 0)

    visit(document)
    return {"warnings": ["Word 提取包含段落和表格；文本框、修订、公式及插图可能未完整识别。"]}


def parse_pptx(path: Path, blocks: Blocks) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    if len(presentation.slides) > LIMITS["max_pages"]:
        fail("幻灯片超过 200 张限制", "ATTACHMENT_TOO_LARGE")

    def shapes(items, number, depth=0):
        if depth > 12:
            fail("幻灯片组合形状层级超过限制")
        for shape in items:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                shapes(shape.shapes, number, depth + 1)
            elif shape.has_table:
                for row in shape.table.rows:
                    blocks.add(
                        " | ".join(cell.text for cell in row.cells), {"slide": number}, "table"
                    )
            elif shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    blocks.add(paragraph.text, {"slide": number}, "slide")

    for number, slide in enumerate(presentation.slides, 1):
        shapes(slide.shapes, number)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            blocks.add(slide.notes_slide.notes_text_frame.text, {"slide": number}, "notes")
    return {
        "pages": len(presentation.slides),
        "warnings": ["已提取幻灯片正文、表格和备注；动画、图表和 SmartArt 的视觉关系可能缺失。"],
    }


def parse_pdf(path: Path, blocks: Blocks) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        fail("暂不支持加密 PDF")
    if len(reader.pages) > LIMITS["max_pages"]:
        fail("PDF 超过 200 页限制", "ATTACHMENT_TOO_LARGE")
    coverage = {"recognized_pages": [], "missing_pages": [], "reasons": {}}
    total_stream = 0
    for number, page in enumerate(reader.pages, 1):
        content = page.get_contents()
        if content is not None:
            total_stream += len(content.get_data())
            if total_stream > ZIP_MAX_BYTES:
                fail("PDF 解压内容超过限制", "ATTACHMENT_TOO_LARGE")
        text = page.extract_text() or ""
        resources = page.get("/Resources") or {}
        resources = resources.get_object() if hasattr(resources, "get_object") else resources
        objects = resources.get("/XObject") or {}
        objects = objects.get_object() if hasattr(objects, "get_object") else objects
        image_only = len(text.strip()) < 30 and any(
            value.get_object().get("/Subtype") == "/Image" for value in objects.values()
        )
        if not text.strip() or image_only:
            coverage["missing_pages"].append(number)
            coverage["reasons"][str(number)] = "本页没有可提取文字，可能是扫描页或空白页"
        else:
            coverage["recognized_pages"].append(number)
            blocks.add(text, {"page": number}, "page")
    if not coverage["recognized_pages"]:
        fail("PDF 没有可提取正文，需要 OCR；首期暂不支持扫描 PDF", "ATTACHMENT_OCR_REQUIRED")
    return {
        "pages": len(reader.pages),
        "coverage": coverage,
        "warnings": ["PDF 的复杂表格、双栏及公式可能无法完整还原。"],
        "status": "needs_confirmation" if coverage["missing_pages"] else "ready",
    }


def convert_legacy(path: Path, extension: str, directory: Path, executable: str | None) -> Path:
    if not executable:
        fail("服务未安装 LibreOffice 转换组件", "ATTACHMENT_CONVERTER_UNAVAILABLE")
    profile = directory / "profile"
    (profile / "user").mkdir(parents=True)
    # Harden the conversion profile even though the whole process has no network access.
    (profile / "user" / "registrymodifications.xcu").write_text(
        """<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>0</value></prop></item></oor:items>"""
    )
    target = "docx" if extension == ".doc" else "pptx"
    result = subprocess.run(
        [
            executable,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            "--convert-to",
            target,
            "--outdir",
            str(directory),
            str(path),
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        timeout=40,
        check=False,
    )
    converted = directory / f"{path.stem}.{target}"
    if result.returncode or not converted.is_file():
        fail("旧版 Office 转换失败，文件可能已加密或损坏")
    if converted.stat().st_size > LIMITS["max_document_bytes"]:
        fail("转换后的文件超过大小限制", "ATTACHMENT_TOO_LARGE")
    validate_file_type(converted, f".{target}")
    return converted


def parse_file(path: Path, extension: str, output: Path, libreoffice: str | None = None) -> dict:
    output.mkdir(parents=True, exist_ok=True)
    validate_file_type(path, extension)
    blocks = Blocks()
    result = {"status": "ready", "warnings": [], "coverage": None, "pages": None}
    original_extension = extension
    if extension in {".doc", ".ppt"}:
        path = convert_legacy(path, extension, output, libreoffice)
        extension = ".docx" if extension == ".doc" else ".pptx"
    if extension in {".jpg", ".jpeg", ".png", ".webp"}:
        from PIL import Image, ImageOps

        with Image.open(path) as source:
            normalized = ImageOps.exif_transpose(source)
            result.update(width=normalized.width, height=normalized.height)
            if normalized.mode != "RGB":
                rgba = normalized.convert("RGBA")
                normalized = Image.new("RGB", rgba.size, "white")
                normalized.paste(rgba, mask=rgba.getchannel("A"))
            normalized.thumbnail((2048, 2048), Image.Resampling.LANCZOS)
            normalized.save(output / "image.jpg", format="JPEG", quality=90)
            normalized.thumbnail((384, 384), Image.Resampling.LANCZOS)
            normalized.save(output / "preview.jpg", format="JPEG", quality=85)
        result.update(image_file="image.jpg", preview_file="preview.jpg")
    elif extension == ".docx":
        result.update(parse_docx(path, blocks))
    elif extension == ".pptx":
        result.update(parse_pptx(path, blocks))
    elif extension == ".pdf":
        result.update(parse_pdf(path, blocks))
    if original_extension in {".doc", ".ppt"}:
        result["warnings"].append(
            "旧版 Office 已转换后提取；引用对应转换后的结构，原始分页和复杂版式可能变化。"
        )
    if extension not in {".jpg", ".jpeg", ".png", ".webp"} and not blocks.items:
        fail("文档没有可提取正文；图片、公式或图表需要另行提供")
    result["chunks"] = blocks.items
    return result


def main() -> None:
    path, extension, directory = Path(sys.argv[1]), sys.argv[2], Path(sys.argv[3])
    import resource

    resource.setrlimit(resource.RLIMIT_CPU, (55, 55))
    resource.setrlimit(resource.RLIMIT_FSIZE, (220 * 1024 * 1024, 220 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_NOFILE, (256, 256))
    # RLIMIT_AS is enforceable on Linux; macOS uses worker RSS sampling instead.
    if sys.platform != "darwin":
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    try:
        result = parse_file(path, extension, directory, sys.argv[4] or None)
    except ParseFailure as exc:
        result = {"error_code": exc.code, "error_message": str(exc)}
    except Exception:
        result = {
            "error_code": "ATTACHMENT_PARSE_FAILED",
            "error_message": "文件损坏、加密或解析失败，请检查后重新上传",
        }
    (directory / "result.json").write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")


if __name__ == "__main__":
    main()

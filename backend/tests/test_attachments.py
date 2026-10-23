from datetime import timedelta
from io import BytesIO
from pathlib import Path
import subprocess
import uuid
import zipfile

import pytest
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from sqlalchemy.orm import sessionmaker

from models import (
    Attachment,
    AttachmentChunk,
    AttachmentCleanup,
    ChatMessage,
    User,
)
from services.attachments import service, worker
from services.attachments.config import LIMITS, isolation_available, libreoffice_path
from services.attachments.parsers import ParseFailure, parse_file, validate_file_type
from services.attachments.storage import AttachmentStorage


@pytest.fixture(autouse=True)
def private_store(tmp_path, monkeypatch):
    monkeypatch.setenv("ATTACHMENT_STORAGE_DIR", str(tmp_path / "private"))


def image_bytes(color="blue"):
    output = BytesIO()
    Image.new("RGB", (64, 32), color).save(output, format="PNG")
    return output.getvalue()


def upload(client, name="picture.png", data=None, **kwargs):
    return client.post(
        "/api/attachments",
        files={"file": (name, data or image_bytes(), "application/octet-stream")},
        data={
            "client_upload_id": str(uuid.uuid4()),
            "session_id": "session-one",
            "mode": "normal",
            **kwargs,
        },
    )


def factory(db):
    return sessionmaker(bind=db.get_bind(), expire_on_commit=False)


def test_upload_idempotency_scope_and_download(client, db_session):
    first = upload(client, client_upload_id="one")
    assert first.status_code == 200, first.text
    item = first.json()
    assert item["status"] == "pending"
    assert upload(client, client_upload_id="one").json()["id"] == item["id"]
    assert upload(client, client_upload_id="one", data=image_bytes("red")).status_code == 409
    assert upload(client, client_upload_id="one", session_id="other").status_code == 409
    response = client.get(f"/api/attachments/{item['id']}/content")
    assert response.content == image_bytes()
    assert "attachment" in response.headers["content-disposition"]
    assert "no-store" in response.headers["cache-control"]
    assert db_session.query(Attachment).count() == 1


def test_unauthorized_every_resource(client, db_session):
    item = upload(client).json()
    other = User(username="other", email="other@test.test", hashed_password="unused")
    db_session.add(other)
    db_session.flush()
    db_session.get(Attachment, item["id"]).user_id = other.id
    db_session.commit()
    for endpoint in ("content", "preview", "chunks/c1"):
        assert client.get(f"/api/attachments/{item['id']}/{endpoint}").status_code == 404
    assert client.get("/api/attachments", params={"ids": item["id"]}).status_code == 404
    assert client.post(f"/api/attachments/{item['id']}/retry").status_code == 404
    assert client.delete(f"/api/attachments/{item['id']}").status_code == 404


def test_type_size_quota_and_batch_limit(client, monkeypatch):
    assert upload(client, name="evil.svg").json()["code"] == "ATTACHMENT_TYPE_UNSUPPORTED"
    assert upload(client, name="false.pdf").json()["code"] == "ATTACHMENT_TYPE_UNSUPPORTED"
    monkeypatch.setitem(LIMITS, "max_image_bytes", 2)
    assert upload(client).status_code == 413
    monkeypatch.setitem(LIMITS, "max_image_bytes", 10 * 1024 * 1024)
    monkeypatch.setitem(LIMITS, "max_user_bytes", 100)
    assert upload(client).json()["code"] == "ATTACHMENT_QUOTA_EXCEEDED"
    assert (
        client.get(
            "/api/attachments", params={"ids": ",".join([str(i) for i in range(51)])}
        ).status_code
        == 422
    )


def test_bound_attachments_immutable_and_delete_after_commit(client, db_session, auth_user):
    item = upload(client).json()
    row = db_session.get(Attachment, item["id"])
    row.status = "ready"
    db_session.commit()
    rows = service.validate_for_message(
        db_session,
        user_id=auth_user.id,
        session_id="session-one",
        topic_id=None,
        mode="normal",
        ids=[row.id],
    )
    message = ChatMessage(
        user_id=auth_user.id,
        session_id="session-one",
        node_key="",
        role="user",
        content="",
        mode="normal",
    )
    db_session.add(message)
    db_session.flush()
    service.bind_to_message(
        db_session, message_id=message.id, session_id="session-one", attachments=rows
    )
    db_session.commit()
    assert service.summaries_for_messages(db_session, [message.id])[message.id][0]["id"] == row.id
    assert client.delete(f"/api/attachments/{row.id}").status_code == 409
    service.delete_for_messages(db_session, [message.id])
    assert AttachmentStorage().path(row.storage_key).is_file()
    db_session.rollback()
    assert client.get(f"/api/attachments/{row.id}/content").status_code == 200
    service.delete_for_messages(db_session, [message.id])
    db_session.commit()
    assert client.get(f"/api/attachments/{row.id}/content").status_code == 404
    assert worker.run_cleanup(factory(db_session))
    assert not AttachmentStorage().path(row.storage_key).exists()


def test_account_cleanup_survives_user_delete(client, db_session, auth_user):
    item = upload(client).json()
    key = db_session.get(Attachment, item["id"]).storage_key
    service.delete_for_user(db_session, auth_user.id)
    db_session.delete(auth_user)
    db_session.commit()
    assert db_session.query(Attachment).count() == 0
    assert db_session.query(AttachmentCleanup).count() == 1
    worker.run_cleanup(factory(db_session))
    assert not AttachmentStorage().path(key).exists()


def pdf(path, *, blank=False, mixed=False):
    writer = PdfWriter()
    page = writer.add_blank_page(300, 300)
    if not blank:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(
            b"BT /F1 14 Tf 30 250 Td (Photosynthesis converts light into energy.) Tj ET"
        )
        page[NameObject("/Contents")] = writer._add_object(stream)
    if mixed:
        writer.add_blank_page(300, 300)
    writer.write(path)


def test_pdf_range_confirmation_and_scan_failure(tmp_path, client, db_session, auth_user):
    source = tmp_path / "mixed.pdf"
    pdf(source, mixed=True)
    parsed = parse_file(source, ".pdf", tmp_path / "out")
    assert parsed["status"] == "needs_confirmation"
    assert parsed["coverage"]["recognized_pages"] == [1]
    assert parsed["coverage"]["missing_pages"] == [2]
    item = upload(client, name="mixed.pdf", data=source.read_bytes()).json()
    worker.run_one_attachment_job(factory(db_session), parser=lambda *_: parsed)
    db_session.expire_all()
    row = db_session.get(Attachment, item["id"])
    assert row.status == "needs_confirmation"
    with pytest.raises(Exception, match="解析完成"):
        service.validate_for_message(
            db_session,
            user_id=auth_user.id,
            session_id="session-one",
            topic_id=None,
            mode="normal",
            ids=[row.id],
        )
    assert (
        client.post(
            f"/api/attachments/{row.id}/accept-text-coverage", json={"revision": 2}
        ).status_code
        == 409
    )
    accepted = client.post(f"/api/attachments/{row.id}/accept-text-coverage", json={"revision": 1})
    assert accepted.json()["status"] == "ready"
    pdf(source, blank=True)
    with pytest.raises(ParseFailure) as failure:
        parse_file(source, ".pdf", tmp_path / "out2")
    assert failure.value.code == "ATTACHMENT_OCR_REQUIRED"


def test_late_parse_cannot_revive_deleted_or_taken_lease(client, db_session):
    item = upload(client).json()

    def deleted_during_parse(*_):
        with factory(db_session)() as db:
            service.mark_deleted(db, [db.get(Attachment, item["id"])])
            db.commit()
        return {"status": "ready", "chunks": [], "warnings": []}

    assert worker.run_one_attachment_job(factory(db_session), parser=deleted_during_parse)
    db_session.expire_all()
    assert db_session.get(Attachment, item["id"]).status == "deleted"
    assert db_session.query(AttachmentChunk).count() == 0


def test_expired_lease_recovered_and_draft_expired(client, db_session):
    item = upload(client).json()
    row = db_session.get(Attachment, item["id"])
    row.status = "processing"
    row.lease_token = "old"
    row.lease_until = service.utcnow() - timedelta(seconds=1)
    db_session.commit()
    assert worker.run_one_attachment_job(
        factory(db_session), parser=lambda *_: {"status": "ready", "chunks": [], "warnings": []}
    )
    db_session.expire_all()
    assert db_session.get(Attachment, item["id"]).status == "ready"
    row = db_session.get(Attachment, item["id"])
    row.expires_at = service.utcnow() - timedelta(seconds=1)
    db_session.commit()
    assert client.get(f"/api/attachments/{item['id']}/content").status_code == 404
    worker.run_cleanup(factory(db_session))
    db_session.expire_all()
    assert db_session.get(Attachment, item["id"]).status == "deleted"


def make_docx(path):
    from docx import Document

    doc = Document()
    doc.add_heading("中文讲义 Photosynthesis", 1)
    doc.add_paragraph("光合作用将光能转化为化学能。")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "输入：光能"
    cell = table.cell(0, 1)
    cell.text = "输出：化学能"
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "第三层嵌套"
    doc.add_paragraph("最终段落")
    doc.save(path)


def make_pptx(path):
    from pptx import Presentation

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "中文课件 Photosynthesis"
    slide.placeholders[1].text = "光合作用将光能转化为化学能。"
    slide.notes_slide.notes_text_frame.text = "教师备注：先询问学生对能量的理解。"
    deck.save(path)


def test_docx_order_nested_tables_pptx_notes(tmp_path):
    doc = tmp_path / "lesson.docx"
    make_docx(doc)
    result = parse_file(doc, ".docx", tmp_path / "out")
    text = "\n".join(chunk["text"] for chunk in result["chunks"])
    assert "第三层嵌套" in text
    assert text.index("输入") < text.index("最终段落")
    assert result["chunks"][0]["locator"]["paragraph"] == 1
    slides = tmp_path / "lesson.pptx"
    make_pptx(slides)
    result = parse_file(slides, ".pptx", tmp_path / "slides")
    assert result["pages"] == 1
    assert any(
        chunk["block_type"] == "notes" and "教师备注" in chunk["text"] for chunk in result["chunks"]
    )


def test_office_archive_bomb_and_path_rejected(tmp_path):
    source = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("../escape", "bad")
    with pytest.raises(ParseFailure, match="路径"):
        validate_file_type(source, ".docx")
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "a" * 1_000_000)
        archive.writestr("[Content_Types].xml", "types")
    with pytest.raises(ParseFailure, match="压缩比例"):
        validate_file_type(source, ".docx")


@pytest.mark.skipif(not isolation_available(), reason="OS sandbox runtime is not installed")
def test_real_sandbox_image_and_document(tmp_path):
    picture = tmp_path / "picture.png"
    picture.write_bytes(image_bytes())
    output = tmp_path / "image-out"
    output.mkdir()
    result = worker.execute_parser(picture, output)
    assert result["status"] == "ready", result
    assert (output / "image.jpg").is_file()
    doc = tmp_path / "lesson.docx"
    make_docx(doc)
    output = tmp_path / "doc-out"
    output.mkdir()
    result = worker.execute_parser(doc, output)
    assert result["status"] == "ready", result
    assert any("光合作用" in chunk["text"] for chunk in result["chunks"])


@pytest.mark.skipif(
    not isolation_available() or not libreoffice_path(),
    reason="LibreOffice and OS sandbox are required",
)
@pytest.mark.parametrize(
    "extension,modern,maker", [("doc", "docx", make_docx), ("ppt", "pptx", make_pptx)]
)
def test_real_legacy_conversion_in_sandbox(tmp_path, extension, modern, maker):
    source = tmp_path / f"lesson.{modern}"
    maker(source)
    profile = tmp_path / "fixture-profile"
    conversion = subprocess.run(
        [
            libreoffice_path(),
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--convert-to",
            extension,
            "--outdir",
            str(tmp_path),
            str(source),
        ],
        capture_output=True,
        timeout=45,
    )
    legacy = tmp_path / f"lesson.{extension}"
    assert conversion.returncode == 0 and legacy.is_file(), conversion.stderr.decode()
    output = tmp_path / "parsed"
    output.mkdir()
    result = worker.execute_parser(legacy, output)
    assert result["status"] == "ready", result
    assert any("光合作用" in chunk["text"] for chunk in result["chunks"])
    if extension == "ppt":
        assert result["pages"] == 1
        assert any(
            chunk["block_type"] == "notes" and "教师备注" in chunk["text"]
            for chunk in result["chunks"]
        )


@pytest.mark.skipif(not isolation_available(), reason="OS sandbox runtime is not installed")
def test_sandbox_denies_other_files_and_network(tmp_path):
    import json
    import os
    import sys

    source = tmp_path / "picture.png"
    source.write_bytes(image_bytes())
    secret = tmp_path / "private-secret.txt"
    secret.write_text("not-readable-by-parser")
    output = tmp_path / "out"
    output.mkdir()
    command = worker._sandbox_command(source, output, None)
    python_index = command.index(sys.executable)
    code = f"""
from pathlib import Path
import socket
blocked = 0
try:
    Path({json.dumps(str(secret))}).read_bytes()
except (PermissionError, FileNotFoundError):
    blocked += 1
try:
    socket.create_connection(("1.1.1.1", 443), timeout=1)
except OSError:
    blocked += 1
print(blocked)
"""
    result = subprocess.run(
        [*command[:python_index], sys.executable, "-I", "-c", code],
        capture_output=True,
        cwd=output,
        env={"PATH": os.environ.get("PATH", "/usr/bin:/bin"), "HOME": str(output)},
        timeout=5,
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == b"2"


def test_slots_limit_jobs_and_stale_token_cannot_publish(client, db_session, tmp_path):
    for _ in range(3):
        assert upload(client).status_code == 200
    session_factory = factory(db_session)
    first = worker._claim(session_factory)
    second = worker._claim(session_factory)
    assert first and second and first["id"] != second["id"]
    assert worker._claim(session_factory) is None
    with session_factory() as db:
        row = db.get(Attachment, first["id"])
        row.lease_token = "reclaimed-by-other-worker"
        db.commit()
    worker._publish_result(
        session_factory, first, {"status": "ready", "chunks": [], "warnings": []}, tmp_path
    )
    db_session.expire_all()
    assert db_session.get(Attachment, first["id"]).status == "processing"
    assert db_session.get(Attachment, first["id"]).lease_token == "reclaimed-by-other-worker"


def test_retry_is_bounded_and_increments_revision(client, db_session):
    item = upload(client).json()
    row = db_session.get(Attachment, item["id"])
    row.status = "failed"
    row.error_code = "ATTACHMENT_PARSE_FAILED"
    row.attempts = 1
    db_session.commit()
    result = client.post(f"/api/attachments/{row.id}/retry")
    assert result.status_code == 200
    assert result.json()["revision"] == 2
    assert result.json()["status"] == "pending"
    row.status = "failed"
    row.attempts = 3
    db_session.commit()
    assert client.post(f"/api/attachments/{row.id}/retry").status_code == 409


def test_attachment_migration_slots_and_capabilities(tmp_path):
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import create_engine, inspect, text

    backend = Path(__file__).resolve().parents[1]
    config = Config(str(backend / "alembic.ini"))
    config.set_main_option("script_location", str(backend / "alembic"))
    database_url = f"sqlite:///{tmp_path / 'migration.sqlite'}"
    config.set_main_option("sqlalchemy.url", database_url)
    command.upgrade(config, "head")
    engine = create_engine(database_url)
    inspector = inspect(engine)
    assert {
        "attachments",
        "attachment_chunks",
        "message_attachments",
        "attachment_cleanup",
        "attachment_worker_slots",
    }.issubset(inspector.get_table_names())
    assert "preparation_json" in {column["name"] for column in inspector.get_columns("chat_turns")}
    assert {"image_input", "max_images", "max_image_bytes", "vision_token_budget"}.issubset(
        {column["name"] for column in inspector.get_columns("user_llm_profiles")}
    )
    with engine.connect() as connection:
        assert connection.execute(
            text("SELECT id FROM attachment_worker_slots ORDER BY id")
        ).scalars().all() == [1, 2]
    command.downgrade(config, "0011_web_search_quotas")
    assert "attachments" not in inspect(engine).get_table_names()
    engine.dispose()


def test_multipart_body_limit_applies_before_spooling(client, monkeypatch):
    monkeypatch.setitem(LIMITS, "max_document_bytes", 128)
    monkeypatch.setitem(LIMITS, "max_image_bytes", 128)
    response = client.post(
        "/api/attachments",
        content=b"x" * 100_000,
        headers={"content-type": "multipart/form-data; boundary=example"},
    )
    assert response.status_code == 413
    assert response.json()["code"] == "ATTACHMENT_TOO_LARGE"


def test_streamed_multipart_limit_does_not_trust_content_length(client, monkeypatch):
    monkeypatch.setitem(LIMITS, "max_document_bytes", 128)
    monkeypatch.setitem(LIMITS, "max_image_bytes", 128)

    def body():
        yield b'--test\r\nContent-Disposition: form-data; name="file"; filename="x.pdf"\r\nContent-Type: application/pdf\r\n\r\n'
        yield b"x" * 100_000
        yield b"\r\n--test--\r\n"

    response = client.post(
        "/api/attachments",
        content=body(),
        headers={"content-type": "multipart/form-data; boundary=test"},
    )
    assert response.status_code == 413
    assert response.json()["code"] == "ATTACHMENT_TOO_LARGE"


def test_encrypted_pdf_and_page_limits(tmp_path, monkeypatch):
    source = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    writer.encrypt("password")
    writer.write(source)
    with pytest.raises(ParseFailure, match="加密"):
        parse_file(source, ".pdf", tmp_path / "out")
    writer = PdfWriter()
    for _ in range(3):
        writer.add_blank_page(100, 100)
    writer.write(source)
    monkeypatch.setitem(LIMITS, "max_pages", 2)
    with pytest.raises(ParseFailure) as failure:
        parse_file(source, ".pdf", tmp_path / "out")
    assert failure.value.code == "ATTACHMENT_TOO_LARGE"


@pytest.mark.skipif(not isolation_available(), reason="OS sandbox runtime is not installed")
def test_parser_deadline_terminates_process(tmp_path, monkeypatch):
    source = tmp_path / "source.png"
    source.write_bytes(image_bytes())
    output = tmp_path / "out"
    output.mkdir()
    monkeypatch.setattr(worker, "PARSER_TIMEOUT_SECONDS", 0)
    result = worker.execute_parser(source, output)
    assert result["error_code"] == "ATTACHMENT_PARSE_FAILED"
    assert "超时" in result["error_message"]
